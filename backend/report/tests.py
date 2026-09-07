import datetime as dt

from django.db import IntegrityError, connection
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from .models import FollowUpWA, Instansi, Kunjungan, User


class BaseTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.spv = User.objects.create_user("devy", password="x", role=User.Role.ADMIN, first_name="Devy")
        cls.sales1 = User.objects.create_user("utoro", password="x", role=User.Role.SALES, first_name="Utoro")
        cls.sales2 = User.objects.create_user("maryadi", password="x", role=User.Role.SALES, first_name="Maryadi")
        cls.inst1 = Instansi.objects.create(nama="SMKN 1 Purworejo", pic="Budi", no_wa="081234567890", alamat="Jl. A")
        cls.inst2 = Instansi.objects.create(nama="SMAN 2 Kutoarjo", pic="Sari", no_wa="0899000111")
        cls.k1 = Kunjungan.objects.create(instansi=cls.inst1, sales=cls.sales1, tanggal=dt.date(2026, 5, 10), catatan="rahasia")
        cls.k2 = Kunjungan.objects.create(instansi=cls.inst2, sales=cls.sales2, tanggal=dt.date(2026, 6, 3), status=Kunjungan.Status.CLOSING)


class AksesRoleTest(BaseTest):
    def test_anonim_diarahkan_ke_login(self):
        r = self.client.get(reverse("dashboard"))
        self.assertEqual(r.status_code, 302)
        self.assertIn(reverse("login"), r["Location"])

    def test_sales_tidak_bisa_lihat_kunjungan_sales_lain(self):
        self.client.force_login(self.sales1)
        self.assertEqual(self.client.get(reverse("kunjungan_detail", args=[self.k2.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("kunjungan_edit", args=[self.k2.pk])).status_code, 404)
        r = self.client.get(reverse("kunjungan_list"))
        self.assertContains(r, "SMKN 1 Purworejo")
        self.assertNotContains(r, "SMAN 2 Kutoarjo")

    def test_spv_lihat_semua_tapi_read_only(self):
        self.client.force_login(self.spv)
        r = self.client.get(reverse("kunjungan_list"))
        self.assertContains(r, "SMKN 1 Purworejo")
        self.assertContains(r, "SMAN 2 Kutoarjo")
        r = self.client.get(reverse("kunjungan_detail", args=[self.k2.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.context["read_only"])
        self.assertNotContains(r, 'data-testid="btn-edit"')
        self.assertEqual(self.client.get(reverse("kunjungan_baru")).status_code, 403)
        self.assertEqual(self.client.get(reverse("kunjungan_edit", args=[self.k2.pk])).status_code, 403)
        r = self.client.post(reverse("klik_wa", args=[self.k2.pk, "H2"]))
        self.assertEqual(r.status_code, 403)

    def test_sales_dashboard_dan_spv_dashboard(self):
        self.client.force_login(self.sales1)
        self.assertTemplateUsed(self.client.get(reverse("dashboard")), "report/dashboard_sales.html")
        self.client.force_login(self.spv)
        self.assertTemplateUsed(self.client.get(reverse("dashboard")), "report/dashboard_spv.html")


class LoginTest(BaseTest):
    def test_login_menampilkan_captcha(self):
        r = self.client.get(reverse("login"))
        self.assertContains(r, 'name="captcha_1"')

    def test_login_sukses_dengan_captcha(self):
        r = self.client.post(
            reverse("login"),
            {"username": "utoro", "password": "x", "captcha_0": "dummy", "captcha_1": "PASSED"},
        )
        self.assertRedirects(r, reverse("dashboard"))

    def test_login_captcha_salah_ditolak(self):
        r = self.client.post(
            reverse("login"),
            {"username": "utoro", "password": "x", "captcha_0": "dummy", "captcha_1": "salah"},
        )
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.wsgi_request.user.is_authenticated)

    def test_logout_hanya_post(self):
        self.client.force_login(self.sales1)
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)
        self.assertRedirects(self.client.post(reverse("logout")), reverse("login"))


class CekInstansiTest(BaseTest):
    def test_sudah_dikunjungi(self):
        self.client.force_login(self.sales2)
        r = self.client.get(reverse("cek_instansi"), {"nama": "smkn 1 purworejo"})
        data = r.json()
        self.assertTrue(data["exists"])
        self.assertTrue(data["sudah_dikunjungi"])
        self.assertEqual(data["oleh"], "Utoro")
        self.assertIn("Sudah dikunjungi oleh Utoro", data["pesan"])
        self.assertEqual(data["pic"], "Budi")
        self.assertEqual(data["no_wa"], "081234567890")

    def test_belum_ada(self):
        self.client.force_login(self.sales1)
        self.assertFalse(self.client.get(reverse("cek_instansi"), {"nama": "Tidak Ada"}).json()["exists"])


class KunjunganFormTest(BaseTest):
    def test_buat_kunjungan_reuse_instansi_case_insensitive(self):
        self.client.force_login(self.sales2)
        r = self.client.post(
            reverse("kunjungan_baru"),
            {"nama_instansi": "smkn 1 PURWOREJO", "tanggal": "2026-06-05", "pic": "Budi", "no_wa": "0812-3456-7890",
             "alamat": "Jl. A", "catatan": "ok", "status": "PEMBICARAAN"},
        )
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Instansi.objects.count(), 2)
        k = Kunjungan.objects.get(sales=self.sales2, instansi=self.inst1)
        self.assertEqual(k.status, "PEMBICARAAN")
        self.inst1.refresh_from_db()
        self.assertEqual(self.inst1.no_wa, "6281234567890")

    def test_edit_kunjungan_sendiri(self):
        self.client.force_login(self.sales1)
        r = self.client.post(
            reverse("kunjungan_edit", args=[self.k1.pk]),
            {"nama_instansi": "SMKN 1 Purworejo", "tanggal": "2026-05-10", "status": "CLOSING", "catatan": "deal"},
        )
        self.assertRedirects(r, reverse("kunjungan_detail", args=[self.k1.pk]))
        self.k1.refresh_from_db()
        self.assertEqual(self.k1.status, "CLOSING")
        self.assertEqual(self.k1.catatan, "deal")


class KlikWATest(BaseTest):
    def test_post_mengubah_tahap(self):
        self.client.force_login(self.sales1)
        r = self.client.post(reverse("klik_wa", args=[self.k1.pk, "H0"]))
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r["Location"].startswith("https://wa.me/6281234567890?text="))
        self.k1.refresh_from_db()
        self.assertEqual(self.k1.tahap, "AWAL")
        self.client.post(reverse("klik_wa", args=[self.k1.pk, "H2"]))
        self.k1.refresh_from_db()
        self.assertEqual(self.k1.tahap, "FOLLOW_UP")
        self.client.post(reverse("klik_wa", args=[self.k1.pk, "H5"]))
        self.k1.refresh_from_db()
        self.assertEqual(self.k1.tahap, "PROGRESS")
        self.assertEqual(FollowUpWA.objects.filter(kunjungan=self.k1).count(), 3)
        self.assertEqual(FollowUpWA.objects.filter(kunjungan=self.k1, jenis="H2").first().oleh, self.sales1)

    def test_get_ditolak(self):
        self.client.force_login(self.sales1)
        self.assertEqual(self.client.get(reverse("klik_wa", args=[self.k1.pk, "H2"])).status_code, 405)

    def test_jenis_tidak_valid(self):
        self.client.force_login(self.sales1)
        self.assertEqual(self.client.post(reverse("klik_wa", args=[self.k1.pk, "H9"])).status_code, 404)

    def test_detail_menampilkan_log_wa(self):
        self.client.force_login(self.sales1)
        self.client.post(reverse("klik_wa", args=[self.k1.pk, "H2"]))
        r = self.client.get(reverse("kunjungan_detail", args=[self.k1.pk]))
        self.assertContains(r, "WA H2 dikirim")
        self.assertContains(r, "rahasia")


class ReminderTest(BaseTest):
    def test_reminder_h2_dan_h5(self):
        today = timezone.localdate()
        k_h2 = Kunjungan.objects.create(instansi=self.inst1, sales=self.sales1, tanggal=today - dt.timedelta(days=2))
        Kunjungan.objects.create(instansi=self.inst2, sales=self.sales1, tanggal=today - dt.timedelta(days=5))
        self.client.force_login(self.sales1)
        r = self.client.get(reverse("blast_wa"))
        self.assertEqual(r.context["reminder_count"], 2)
        self.assertContains(r, f'data-testid="badge-reminder-{k_h2.pk}"')
        FollowUpWA.objects.create(kunjungan=k_h2, jenis="H2", oleh=self.sales1)
        r = self.client.get(reverse("blast_wa"))
        self.assertEqual(r.context["reminder_count"], 1)
        self.client.force_login(self.sales2)
        self.assertEqual(self.client.get(reverse("blast_wa")).context["reminder_count"], 0)


class InstansiUnikTest(TestCase):
    def test_unik_case_insensitive(self):
        Instansi.objects.create(nama="SMKN 1")
        with self.assertRaises(IntegrityError):
            Instansi.objects.create(nama="smkn 1")


class EnkripsiTest(TestCase):
    def test_tersimpan_terenkripsi_dan_terbaca(self):
        inst = Instansi.objects.create(nama="Enkrip", pic="Pak Joko", no_wa="0811", alamat="Jl. Rahasia 1")
        k = Kunjungan.objects.create(instansi=inst, sales=User.objects.create_user("s", password="x"), catatan="catatan rahasia")
        with connection.cursor() as cur:
            cur.execute("SELECT pic, no_wa, alamat FROM report_instansi WHERE id = %s", [inst.pk])
            pic, no_wa, alamat = cur.fetchone()
            cur.execute("SELECT catatan FROM report_kunjungan WHERE id = %s", [k.pk])
            (catatan,) = cur.fetchone()
        for raw, plain in ((pic, "Pak Joko"), (no_wa, "0811"), (alamat, "Jl. Rahasia 1"), (catatan, "catatan rahasia")):
            self.assertNotEqual(raw, plain)
            self.assertNotIn(plain, raw)
        inst.refresh_from_db()
        k.refresh_from_db()
        self.assertEqual((inst.pic, inst.no_wa, inst.alamat, k.catatan), ("Pak Joko", "0811", "Jl. Rahasia 1", "catatan rahasia"))

    def test_hard_fail_tanpa_key_saat_debug_off(self):
        from django.core.exceptions import ImproperlyConfigured
        from django.test import override_settings

        from .crypto_fields import _get_fernet

        with override_settings(DEBUG=False, FIELD_ENCRYPTION_KEY=""):
            with self.assertRaises(ImproperlyConfigured):
                _get_fernet()


class ExportExcelTest(BaseTest):
    def test_export(self):
        self.client.force_login(self.spv)
        r = self.client.get(reverse("export_excel"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn(".xlsx", r["Content-Disposition"])
        self.assertTrue(r.content.startswith(b"PK"))

    def test_export_sales_hanya_data_sendiri(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.force_login(self.sales1)
        r = self.client.get(reverse("export_excel"))
        ws = load_workbook(BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], "SMKN 1 Purworejo")


class DashboardFilterTest(BaseTest):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        Kunjungan.objects.create(instansi=cls.inst1, sales=cls.sales2, tanggal=dt.date(2026, 6, 10), status="CLOSING")
        Kunjungan.objects.create(instansi=cls.inst2, sales=cls.sales1, tanggal=dt.date(2026, 6, 12), status="BARU")

    def setUp(self):
        self.client.force_login(self.spv)

    def stats(self, **params):
        return self.client.get(reverse("dashboard"), params).context["stats"]

    def test_tanpa_filter(self):
        s = self.stats()
        self.assertEqual((s["total"], s["instansi"], s["closing"]), (4, 2, 2))

    def test_filter_bulan(self):
        s = self.stats(bulan="2026-06")
        self.assertEqual((s["total"], s["closing"]), (3, 2))
        self.assertEqual(self.stats(bulan="2026-05")["total"], 1)

    def test_filter_sales(self):
        s = self.stats(sales=self.sales2.pk)
        self.assertEqual((s["total"], s["closing"]), (2, 2))

    def test_filter_status_dan_kombinasi(self):
        self.assertEqual(self.stats(status="BARU")["total"], 2)
        s = self.stats(bulan="2026-06", sales=self.sales1.pk, status="BARU")
        self.assertEqual(s["total"], 1)

    def test_grafik_per_sales(self):
        r = self.client.get(reverse("dashboard"))
        per = {p["nama"]: p["n"] for p in r.context["per_sales"]}
        self.assertEqual(per, {"Utoro": 2, "Maryadi": 2})
