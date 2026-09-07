import datetime as dt
from functools import wraps

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Count
from django.http import Http404, HttpResponse, HttpResponseRedirect, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.formats import date_format
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .forms import KunjunganForm, LoginForm
from .models import FollowUpWA, Instansi, Kunjungan, User
from .wa_templates import TAHAP_SETELAH, buat_pesan, link_wa

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def kunjungan_untuk(user):
    qs = Kunjungan.objects.select_related("instansi", "sales")
    return qs if user.is_spv else qs.filter(sales=user)


def sales_only(view):
    @wraps(view)
    def inner(request, *args, **kwargs):
        if not request.user.is_sales:
            raise PermissionDenied("Hanya Sales yang dapat mengubah data kunjungan.")
        return view(request, *args, **kwargs)

    return login_required(inner)


def terapkan_filter(qs, params):
    bulan = params.get("bulan", "").strip()
    sales = params.get("sales", "").strip()
    status = params.get("status", "").strip()
    if bulan:
        try:
            tahun, bln = (int(x) for x in bulan.split("-"))
            qs = qs.filter(tanggal__year=tahun, tanggal__month=bln)
        except ValueError:
            bulan = ""
    if sales.isdigit():
        qs = qs.filter(sales_id=int(sales))
    else:
        sales = ""
    if status in Kunjungan.Status.values:
        qs = qs.filter(status=status)
    else:
        status = ""
    return qs, {"bulan": bulan, "sales": sales, "status": status}


def hitung_reminder(user):
    if not user.is_authenticated:
        return 0
    qs = (
        kunjungan_untuk(user)
        .filter(tanggal__in=Kunjungan.tanggal_reminder())
        .exclude(status=Kunjungan.Status.BATAL)
        .prefetch_related("followup")
    )
    return sum(1 for k in qs if k.reminder_hari_ini())


def root(request):
    return redirect("dashboard")


def login_view(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    form = LoginForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = authenticate(
            request,
            username=form.cleaned_data["username"],
            password=form.cleaned_data["password"],
        )
        if user is None:
            form.add_error(None, "Username atau password salah.")
        else:
            login(request, user)
            nxt = request.POST.get("next") or request.GET.get("next")
            if nxt and url_has_allowed_host_and_scheme(nxt, {request.get_host()}, request.is_secure()):
                return redirect(nxt)
            return redirect("dashboard")
    return render(request, "report/login.html", {"form": form, "next": request.GET.get("next", "")})


@require_POST
def logout_view(request):
    logout(request)
    return redirect("login")


@login_required
def dashboard(request):
    if request.user.is_spv:
        return _dashboard_spv(request)
    return _dashboard_sales(request)


def _bulan_opsi():
    today = timezone.localdate().replace(day=1)
    opsi = []
    for i in range(12):
        tahun = today.year + (today.month - 1 - i) // 12
        bln = (today.month - 1 - i) % 12 + 1
        d = dt.date(tahun, bln, 1)
        opsi.append((f"{tahun}-{bln:02d}", date_format(d, "F Y")))
    return opsi


def _dashboard_spv(request):
    qs, f = terapkan_filter(Kunjungan.objects.select_related("instansi", "sales"), request.GET)
    today = timezone.localdate()
    stats = {
        "total": qs.count(),
        "instansi": qs.values("instansi").distinct().count(),
        "closing": qs.filter(status=Kunjungan.Status.CLOSING).count(),
        "bulan_ini": qs.filter(tanggal__year=today.year, tanggal__month=today.month).count(),
    }
    per_sales = list(
        qs.values("sales__id", "sales__first_name", "sales__last_name", "sales__username")
        .annotate(n=Count("id"))
        .order_by("-n", "sales__first_name")
    )
    maks = max((p["n"] for p in per_sales), default=0)
    for p in per_sales:
        p["nama"] = f"{p['sales__first_name']} {p['sales__last_name']}".strip() or p["sales__username"]
        p["pct"] = round(p["n"] * 100 / maks) if maks else 0
    return render(
        request,
        "report/dashboard_spv.html",
        {
            "stats": stats,
            "per_sales": per_sales,
            "kunjungan": qs[:100],
            "filter": f,
            "bulan_opsi": _bulan_opsi(),
            "sales_opsi": User.objects.filter(role=User.Role.SALES, is_active=True).order_by("first_name"),
            "status_opsi": Kunjungan.Status.choices,
            "query": request.GET.urlencode(),
        },
    )


def _dashboard_sales(request):
    qs = kunjungan_untuk(request.user)
    today = timezone.localdate()
    stats = {
        "total": qs.count(),
        "bulan_ini": qs.filter(tanggal__year=today.year, tanggal__month=today.month).count(),
        "closing": qs.filter(status=Kunjungan.Status.CLOSING).count(),
        "reminder": hitung_reminder(request.user),
    }
    return render(request, "report/dashboard_sales.html", {"stats": stats, "terbaru": qs[:5]})


@login_required
def kunjungan_list(request):
    qs = kunjungan_untuk(request.user)
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(instansi__nama__icontains=q)
    return render(request, "report/kunjungan_list.html", {"kunjungan": qs[:200], "q": q})


def _render_form(request, form, judul, kunjungan=None):
    return render(
        request,
        "report/kunjungan_form.html",
        {
            "form": form,
            "judul": judul,
            "kunjungan": kunjungan,
            "instansi_names": list(Instansi.objects.values_list("nama", flat=True)),
        },
    )


@sales_only
def kunjungan_baru(request):
    form = KunjunganForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        k = form.simpan(request.user)
        messages.success(request, f"Kunjungan ke {k.instansi.nama} tersimpan.")
        return redirect("kunjungan_detail", pk=k.pk)
    return _render_form(request, form, "Kunjungan Baru")


@sales_only
def kunjungan_edit(request, pk):
    k = get_object_or_404(Kunjungan.objects.select_related("instansi"), pk=pk, sales=request.user)
    form = KunjunganForm(request.POST or None, instance=k)
    if request.method == "POST" and form.is_valid():
        form.simpan(request.user)
        messages.success(request, "Kunjungan diperbarui.")
        return redirect("kunjungan_detail", pk=k.pk)
    return _render_form(request, form, "Edit Kunjungan", k)


@login_required
def cek_instansi(request):
    nama = " ".join(request.GET.get("nama", "").split())
    if not nama:
        return JsonResponse({"exists": False})
    instansi = Instansi.objects.filter(nama__iexact=nama).first()
    if instansi is None:
        return JsonResponse({"exists": False})
    terakhir = instansi.kunjungan.select_related("sales")
    exclude = request.GET.get("exclude", "")
    if exclude.isdigit():
        terakhir = terakhir.exclude(pk=int(exclude))
    terakhir = terakhir.first()
    data = {
        "exists": True,
        "nama": instansi.nama,
        "alamat": instansi.alamat,
        "pic": instansi.pic,
        "no_wa": instansi.no_wa,
        "sudah_dikunjungi": terakhir is not None,
    }
    if terakhir:
        tgl = date_format(terakhir.tanggal, "d M Y")
        data.update(
            oleh=terakhir.sales.nama,
            tanggal=tgl,
            status=terakhir.get_status_display(),
            pesan=f"Sudah dikunjungi oleh {terakhir.sales.nama}, {tgl}",
        )
    return JsonResponse(data)


@login_required
def blast_wa(request):
    rows = list(
        kunjungan_untuk(request.user)
        .exclude(status=Kunjungan.Status.BATAL)
        .prefetch_related("followup")[:200]
    )
    for k in rows:
        k.terkirim = k.jenis_terkirim()
        k.reminder = k.reminder_hari_ini()
    return render(
        request,
        "report/blast_wa.html",
        {"kunjungan": rows, "jenis_list": FollowUpWA.Jenis.values, "bisa_kirim": request.user.is_sales},
    )


@require_POST
@sales_only
def klik_wa(request, pk, jenis):
    if jenis not in FollowUpWA.Jenis.values:
        raise Http404
    k = get_object_or_404(kunjungan_untuk(request.user), pk=pk)
    if not k.instansi.no_wa:
        messages.error(request, f"{k.instansi.nama} belum memiliki nomor WhatsApp.")
        return redirect("blast_wa")
    pesan = buat_pesan(jenis, k)
    FollowUpWA.objects.create(kunjungan=k, jenis=jenis, oleh=request.user, pesan=pesan)
    tahap = TAHAP_SETELAH.get(jenis)
    if tahap and k.tahap != tahap:
        k.tahap = tahap
        k.save(update_fields=["tahap", "diubah_pada"])
    return HttpResponseRedirect(link_wa(k.instansi.no_wa, pesan))


@login_required
def kunjungan_detail(request, pk):
    k = get_object_or_404(kunjungan_untuk(request.user).prefetch_related("followup__oleh"), pk=pk)
    timeline = [
        {
            "waktu": date_format(k.tanggal, "d M Y"),
            "judul": "Kunjungan dicatat",
            "keterangan": f"Status {k.get_status_display()} oleh {k.sales.nama}",
            "tipe": "kunjungan",
        }
    ]
    for f in k.followup.all():
        timeline.append(
            {
                "waktu": date_format(timezone.localtime(f.waktu), "d M Y H:i"),
                "judul": f"WA {f.jenis} dikirim",
                "keterangan": f"oleh {f.oleh.nama if f.oleh else '-'}",
                "tipe": "wa",
            }
        )
    if k.catatan:
        timeline.append({"waktu": "Catatan", "judul": "Catatan kunjungan", "keterangan": k.catatan, "tipe": "catatan"})
    read_only = not (request.user.is_sales and k.sales_id == request.user.id)
    return render(
        request,
        "report/kunjungan_detail.html",
        {"k": k, "timeline": timeline, "read_only": read_only},
    )


@login_required
def export_excel(request):
    qs, _ = terapkan_filter(kunjungan_untuk(request.user), request.GET)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kunjungan"
    header = ["Tanggal", "Instansi", "Sales", "PIC", "No. WA", "Alamat", "Status", "Tahap", "Catatan"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8B1A2B")
    for k in qs:
        ws.append(
            [
                k.tanggal.strftime("%d/%m/%Y"),
                k.instansi.nama,
                k.sales.nama,
                k.instansi.pic,
                k.instansi.no_wa,
                k.instansi.alamat,
                k.get_status_display(),
                k.get_tahap_display(),
                k.catatan,
            ]
        )
    for col, width in zip("ABCDEFGHI", (12, 30, 20, 20, 16, 36, 14, 16, 40)):
        ws.column_dimensions[col].width = width
    response = HttpResponse(content_type=XLSX)
    response["Content-Disposition"] = (
        f'attachment; filename="kunjungan-{timezone.localdate():%Y%m%d}.xlsx"'
    )
    wb.save(response)
    return response
